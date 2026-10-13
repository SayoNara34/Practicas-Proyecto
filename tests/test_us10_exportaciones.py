"""US10 - Archivos válidos, filtros, permisos y descarga HTTP."""
from io import BytesIO
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app import estadisticas, exportaciones, reportes
from app.errores import ErrorNoEncontrado, ErrorPermiso, ErrorValidacion


@pytest.mark.parametrize("formato", ["pdf", "xlsx"])
@pytest.mark.parametrize("zona,plaga,total", [(None, None, 2), ("acacías", None, 1),
                                              ("Acacías", "SIGATOKA", 1), ("Bogotá", None, 0)])
def test_archivo_refleja_filtros_y_totales(conn, campesino, agronomo, formato, zona, plaga, total):
    reportes.crear_reporte(conn, campesino, "Manchas", "Finca", ["a.jpg"], "Sigatoka", "Acacías")
    reportes.crear_reporte(conn, campesino, "Tallos", "Finca", ["b.jpg"], "Picudo", "Villavicencio")
    datos = exportaciones.exportar_estadisticas(conn, agronomo, formato, zona, plaga)
    esperado = estadisticas.reportes_por_estado(conn, agronomo, zona, plaga)
    assert sum(esperado.values()) == total
    if formato == "xlsx":
        hoja = load_workbook(BytesIO(datos)).active
        assert dict(hoja.iter_rows(min_row=7, max_row=10, values_only=True)) == esperado
        assert hoja["B11"].value == total
        assert hoja["B3"].value == (zona or "Todas")
    else:
        texto = " ".join(p.extract_text() for p in PdfReader(BytesIO(datos)).pages)
        assert "Total\n" + str(total) in texto
        for estado, cantidad in esperado.items():
            assert f"{estado}\n{cantidad}" in texto


@pytest.mark.parametrize("formato", ["pdf", "xlsx"])
def test_sin_reportes_genera_archivo_valido(conn, agronomo, formato):
    datos = exportaciones.exportar_estadisticas(conn, agronomo, formato)
    if formato == "pdf":
        assert len(PdfReader(BytesIO(datos)).pages) == 1
    else:
        assert load_workbook(BytesIO(datos)).active["B11"].value == 0


@pytest.mark.parametrize("formato", ["pdf", "xlsx"])
def test_campesino_no_exporta(conn, campesino, formato):
    with pytest.raises(ErrorPermiso):
        exportaciones.exportar_estadisticas(conn, campesino, formato)


def test_usuario_inexistente_no_exporta(conn):
    with pytest.raises(ErrorNoEncontrado):
        exportaciones.exportar_estadisticas(conn, 999, "pdf")


def test_rechaza_formato_invalido(conn, agronomo):
    with pytest.raises(ErrorValidacion):
        exportaciones.exportar_estadisticas(conn, agronomo, "csv")


def test_filtro_excel_permanece_texto(conn, agronomo):
    datos = exportaciones.exportar_estadisticas(conn, agronomo, "xlsx", zona="=1+1", tipo_plaga="<b>&")
    hoja = load_workbook(BytesIO(datos)).active
    assert hoja["B3"].data_type == "s" and hoja["B3"].value == "=1+1"
    assert hoja["B4"].value == "<b>&"


@pytest.mark.parametrize("formato", ["pdf", "xlsx"])
def test_rechaza_filtros_demasiado_largos(conn, agronomo, formato):
    with pytest.raises(ErrorValidacion):
        exportaciones.exportar_estadisticas(conn, agronomo, formato, zona="a" * 201)


def test_descarga_http_y_rechazo_de_acceso(monkeypatch, conn, campesino, agronomo):
    import demo_web
    from urllib.parse import urlencode
    monkeypatch.setattr(demo_web, "CONN", conn)
    for uid, formato, codigo in [(agronomo, "pdf", 200), (agronomo, "xlsx", 200),
                                (campesino, "pdf", 400), (999, "pdf", 400),
                                (agronomo, "csv", 400), ("no-numero", "pdf", 400)]:
        manejador = object.__new__(demo_web.Manejador)
        manejador.path = "/exportar?" + urlencode(dict(u=uid, formato=formato))
        manejador.wfile = BytesIO()
        respuesta = {}
        manejador.send_response = lambda c: respuesta.update(codigo=c)
        manejador.send_header = lambda k, v: respuesta.update({k: v})
        manejador.end_headers = lambda: None
        manejador.do_GET()
        assert respuesta["codigo"] == codigo
        if codigo == 200:
            assert respuesta["Content-Type"] == exportaciones.FORMATOS[formato][0]
            assert "attachment" in respuesta["Content-Disposition"]
            assert int(respuesta["Content-Length"]) == len(manejador.wfile.getvalue())
        else:
            assert respuesta["Content-Type"].startswith("text/html")
