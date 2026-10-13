"""US10 - Exportación de las estadísticas a PDF y Excel para profesionales."""
import sqlite3
from datetime import datetime, timezone
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from . import estadisticas
from .errores import ErrorValidacion

FORMATOS = {
    "pdf": ("application/pdf", "estadisticas.pdf"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "estadisticas.xlsx"),
}


def exportar_estadisticas(conn: sqlite3.Connection, profesional_id: int,
                         formato: str, zona: str | None = None,
                         tipo_plaga: str | None = None) -> bytes:
    """Genera un archivo con los mismos conteos y filtros del panel (US10).

    Args:
        conn: Conexión a SQLite.
        profesional_id: Usuario que solicita la descarga.
        formato: pdf o xlsx.
        zona: Zona exacta opcional, sin distinguir mayúsculas.
        tipo_plaga: Tipo de plaga exacto opcional, sin distinguir mayúsculas.

    Returns:
        Contenido binario del archivo; no escribe ni modifica la base de datos.

    Raises:
        ErrorPermiso: El usuario no es profesional.
        ErrorNoEncontrado: El usuario no existe.
        ErrorValidacion: El formato no está soportado o un filtro es demasiado largo.
    """
    conteo = estadisticas.reportes_por_estado(conn, profesional_id, zona, tipo_plaga)
    if formato not in FORMATOS:
        raise ErrorValidacion("Seleccione un formato válido: PDF o Excel (.xlsx).")
    zona, tipo_plaga = (zona or "").strip(), (tipo_plaga or "").strip()
    if len(zona) > 200 or len(tipo_plaga) > 200:
        raise ErrorValidacion("Los filtros deben tener como máximo 200 caracteres.")
    fecha = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    filas = [["Estado", "Cantidad"], *[[e, t] for e, t in conteo.items()],
             ["Total", sum(conteo.values())]]
    salida = BytesIO()
    if formato == "xlsx":
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Estadísticas"
        for fila in [["Control de Plagas en Plátano"], ["Generado", fecha],
                     ["Zona", zona or "Todas"], ["Tipo de plaga", tipo_plaga or "Todas"], [], *filas]:
            hoja.append(fila)
        # Los filtros son texto del usuario: nunca deben interpretarse como fórmulas.
        for referencia in ("B3", "B4"):
            hoja[referencia].data_type = "s"
        for celda in hoja[6]:
            celda.font = Font(color="FFFFFF", bold=True)
            celda.fill = PatternFill("solid", fgColor="2E6B2E")
        hoja.column_dimensions["A"].width = 36
        hoja.column_dimensions["B"].width = 48
        hoja.freeze_panes = "A7"
        libro.save(salida)
    else:
        estilos = getSampleStyleSheet()
        tabla = Table(filas, colWidths=[320, 120])
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2e6b2e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), .5, colors.lightgrey),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ]))
        contenido = [Paragraph("Control de Plagas en Plátano", estilos["Title"]),
                     Paragraph("Estadísticas de reportes", estilos["Heading2"])]
        for texto in [f"Generado: {fecha}", f"Zona: {zona or 'Todas'}",
                      f"Tipo de plaga: {tipo_plaga or 'Todas'}"]:
            contenido.append(Paragraph(escape(texto), estilos["Normal"]))
        contenido.extend([Spacer(1, 16), tabla])
        SimpleDocTemplate(salida, title="Estadísticas de reportes").build(contenido)
    return salida.getvalue()
